"""Logic tests for the pipeline fixes from the 2026-10-05 audit. No network.

Runs under pytest or directly:  python tests/test_pipeline_logic.py
"""

import sys
import tempfile
from datetime import datetime
from pathlib import Path

import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from src import config, cpi, download, excel_output, main  # noqa: E402
from src.analyse import analyse_month  # noqa: E402


def _month_frame(year: int, month: int, price=50.0, drop=0, nan=0) -> pd.DataFrame:
    """A synthetic AEMO month: exactly the expected interval END stamps, flat price."""
    n = config.expected_interval_count(year, month)
    step = pd.Timedelta("5min") if datetime(year, month, 1) >= config.FORMAT_CHANGE_DATE else pd.Timedelta("30min")
    start = pd.Timestamp(year, month, 1)
    ends = pd.date_range(start + step, periods=n, freq=step)
    assert ends[-1] == (start + pd.offsets.MonthBegin(1)), "synthetic month must end 00:00 on the 1st"
    df = pd.DataFrame({"SETTLEMENTDATE": ends, "RRP": price, "REGION": "NSW1"})
    if nan:
        df.loc[: nan - 1, "RRP"] = float("nan")
    if drop:
        df = df.iloc[:-drop]
    return df


# ----------------------------------------------------------------------------- completeness
def test_expected_interval_counts_at_the_format_change():
    assert config.expected_interval_count(2021, 9) == 30 * 48        # last 30-min month
    assert config.expected_interval_count(2021, 10) == 31 * 288      # first 5-min month


def test_expected_interval_counts_by_era():
    assert config.expected_interval_count(2020, 2) == 29 * 48        # leap year, 30-min era
    assert config.expected_interval_count(2024, 2) == 29 * 288       # leap year, 5-min era
    assert config.expected_interval_count(2026, 9) == 30 * 288


def test_complete_month_is_accepted_and_partial_is_rejected():
    assert analyse_month(_month_frame(2020, 3), "NSW1", 2020, 3)["total_intervals"] == 31 * 48
    assert analyse_month(_month_frame(2020, 3, drop=1), "NSW1", 2020, 3) is None   # one interval short
    assert analyse_month(_month_frame(2026, 9, drop=288 * 3), "NSW1", 2026, 9) is None  # 3 days short


def test_blank_prices_are_not_counted_so_the_month_is_rejected():
    # The old code counted NaN rows in total_intervals (while the mean skipped them), so a month
    # full of blanks still "had" every interval.
    assert analyse_month(_month_frame(2020, 3, nan=5), "NSW1", 2020, 3) is None


# ----------------------------------------------------------------------------- guard
def _summary(rows):
    return pd.DataFrame(rows, columns=["region", "year_month", "rrp_nominal", "peak_rrp_nominal",
                                       "total_intervals", "peak_intervals", "carbon_flag"])


BASE = [
    ("NSW1", "2020-04", 50.0, 55.0, 1440, 660, False),
    ("QLD1", "2020-04", 40.0, 45.0, 1440, 660, False),
    ("NSW1", "2020-05", 51.0, 56.0, 1488, 690, False),
    ("QLD1", "2020-05", 41.0, 46.0, 1488, 690, False),
    ("NSW1", "2020-06", 52.0, 57.0, 1440, 660, False),
    ("QLD1", "2020-06", 42.0, 47.0, 1440, 660, False),
]
MUTABLE = {"2020-06"}


def _raises(fn):
    try:
        fn()
    except RuntimeError:
        return True
    return False


def test_guard_allows_an_unchanged_series_and_a_mutable_month_change():
    before = _summary(BASE)
    after = _summary(BASE[:-2] + [("NSW1", "2020-06", 99.0, 99.0, 1440, 660, False),
                                   ("QLD1", "2020-06", 98.0, 98.0, 1440, 660, False)])
    main._assert_settled_history_unchanged(before, after, MUTABLE)


def test_guard_allows_a_missing_settled_row_to_heal():
    # Last run missed QLD1 2020-05 (download failed). Retrying it now is not "changing history".
    before = _summary([r for r in BASE if r[:2] != ("QLD1", "2020-05")])
    after = _summary(BASE)
    main._assert_settled_history_unchanged(before, after, MUTABLE)   # the old guard raised here


def test_guard_still_rejects_a_changed_settled_value():
    bad = [("NSW1", "2020-05", 51.01, 56.0, 1488, 690, False) if r[:2] == ("NSW1", "2020-05") else r for r in BASE]
    assert _raises(lambda: main._assert_settled_history_unchanged(_summary(BASE), _summary(bad), MUTABLE))


def test_guard_rejects_a_vanished_settled_row():
    after = _summary([r for r in BASE if r[:2] != ("NSW1", "2020-04")])
    assert _raises(lambda: main._assert_settled_history_unchanged(_summary(BASE), after, MUTABLE))


# ----------------------------------------------------------------------------- latest-month probe
class _Resp:
    def __init__(self, status):
        self.status_code = status


def _probe(statuses_by_ym, now, sequence=None):
    """Run get_latest_available_month with a fake clock and fake HEAD responses.

    statuses_by_ym answers by the month in the URL; `sequence` answers by call order instead
    (independent of which months get probed, so it also exercises code that ignores the fake clock).
    """
    seen = []
    real_head, real_now, real_sleep = download.requests.head, config.nem_now, download.time.sleep

    def fake_head(url, **_):
        ym = url.split("PRICE_AND_DEMAND_")[1][:6]
        seen.append(ym)
        s = sequence[min(len(seen), len(sequence)) - 1] if sequence else statuses_by_ym.get(ym, 404)
        if s == "boom":
            raise download.requests.RequestException("down")
        return _Resp(s)

    download.requests.head, config.nem_now, download.time.sleep = fake_head, (lambda: now), (lambda *_: None)
    try:
        return download.get_latest_available_month(), seen
    finally:
        download.requests.head, config.nem_now, download.time.sleep = real_head, real_now, real_sleep


def test_probe_steps_back_by_calendar_month_not_30_days():
    # 1 March: the old `now - 30 days` probe went 1 Mar -> 30 Jan and never looked at February.
    got, seen = _probe({"202602": 200}, datetime(2026, 3, 1, 9, 0))
    assert got == (2026, 2), (got, seen)
    assert seen[:2] == ["202603", "202602"], seen


def test_probe_does_not_answer_a_server_error_with_an_older_month():
    # Every attempt on the first month fails with a 5xx (or a network error); a later month would
    # answer 200. The old code gave up on the month after one 5xx and returned the OLDER month.
    n = config.MAX_RETRIES
    got, seen = _probe({}, datetime(2026, 3, 15), sequence=[503] * n + [200])
    assert got is None, (got, seen)
    got, _ = _probe({}, datetime(2026, 3, 15), sequence=["boom"] * n + [200])
    assert got is None
    # ...but a transient 5xx that recovers within the retries is simply retried.
    got, _ = _probe({}, datetime(2026, 3, 15), sequence=[503, 200])
    assert got == (2026, 3), got


def test_probe_year_boundary():
    got, _ = _probe({"202512": 200}, datetime(2026, 1, 2))
    assert got == (2025, 12)


# ----------------------------------------------------------------------------- CPI parsing
G1 = """G1 CONSUMER PRICE INFLATION
Title,Consumer price index,Year-ended inflation
Description,Consumer price index; All groups,Year-ended
Frequency,Quarterly,Quarterly
Type,Original,Original
Units,"Index, September 2025 month = 100",Per cent

Source,ABS / RBA,ABS / RBA
Publication date,30-Jul-2026,30-Jul-2026
Series ID,GCPIAG,GCPIAGYP
30/06/1922,1.93
30/09/1922,1.93,,
31/12/1922,1.91,,
31/03/1923,,,
"""


def _parse(text):
    with tempfile.TemporaryDirectory() as d:
        p = Path(d) / "g1.csv"
        p.write_text(text)
        return cpi._parse_cpi(p)


def test_cpi_parse_keeps_the_first_data_row_and_drops_blank_future_rows():
    out = _parse(G1)
    assert len(out) == 3, out            # the old fixed skiprows+header=0 swallowed 30/06/1922
    assert out["date"].iloc[0] == pd.Timestamp(1922, 6, 30)
    assert out["cpi_index"].tolist() == [1.93, 1.93, 1.91]


def test_cpi_parse_survives_extra_metadata_rows_and_fails_loudly_without_series_id():
    assert len(_parse("Extra row,1\nAnother,2\n" + G1)) == 3       # a longer metadata block still parses
    try:
        _parse(G1.replace("Series ID", "Something else"))
    except ValueError as exc:
        assert "Series ID" in str(exc)
    else:
        raise AssertionError("a layout change must fail loudly")


def _parse_fails(text, needle):
    try:
        _parse(text)
    except ValueError as exc:
        assert needle in str(exc), exc
    else:
        raise AssertionError(f"expected a ValueError mentioning {needle!r}")


def test_cpi_parse_requires_gcpiag_in_column_b():
    # A reordered table would put a percent-change series in column B and deflate every region by it.
    _parse_fails(G1.replace("Series ID,GCPIAG,GCPIAGYP", "Series ID,GCPIAGYP,GCPIAG"), "GCPIAG")


def test_cpi_parse_requires_quarterly_frequency():
    _parse_fails(G1.replace("Frequency,Quarterly,", "Frequency,Monthly,"), "Quarterly")
    _parse_fails(G1.replace("Frequency,Quarterly,Quarterly\n", ""), "Quarterly")       # row missing


def test_cpi_parse_tolerates_a_byte_order_mark():
    assert len(_parse("\ufeff" + G1)) == 3


def _stale(quarter_end, now):
    try:
        cpi.check_cpi_fresh(pd.Timestamp(quarter_end), now)
    except RuntimeError:
        return True
    return False


def test_cpi_freshness_allows_the_normal_wait_for_a_quarter():
    # Jun qtr is the newest until the Sep qtr lands ~28 Oct; the 1 Nov run may still see only Jun.
    assert not _stale("2026-06-30", datetime(2026, 10, 7))
    assert not _stale("2026-06-30", datetime(2026, 11, 1, 10, 38))
    assert not _stale("2026-06-30", datetime(2026, 11, 30))      # a release a month late: still OK
    # Dec qtr, released late Jan (4th Wednesday from Feb 2027), read on 1 May before the Mar qtr.
    assert not _stale("2026-12-31", datetime(2027, 5, 1))


def test_cpi_freshness_fails_when_g1_stops_updating():
    assert _stale("2026-06-30", datetime(2026, 12, 1, 10, 38))   # Sep qtr more than a month late
    assert _stale("2026-03-31", datetime(2026, 10, 7))           # a whole quarter missed
    assert _stale("2025-06-30", datetime(2026, 10, 7))           # frozen for a year


# ----------------------------------------------------------------------------- workbook
def _frame(months, estimated_last=0):
    ym = [p.strftime("%Y-%m") for p in pd.period_range(end="2026-09", periods=months, freq="M")]
    df = pd.DataFrame({"region": "TAS1", "year_month": ym, "rrp_nominal": 50.0, "peak_rrp_nominal": 55.0,
                       "rrp_real": 60.0, "peak_rrp_real": 65.0, "total_intervals": 1, "peak_intervals": 1,
                       "carbon_flag": False, "cpi_estimated": False})
    if estimated_last:
        df.loc[months - estimated_last:, "cpi_estimated"] = True
    return df


def test_workbook_summary_has_explicit_na_rows_and_real_dollar_label():
    from openpyxl import load_workbook
    with tempfile.TemporaryDirectory() as d:
        path = Path(d) / "TAS.xlsx"
        excel_output._write_region_workbook(_frame(150, estimated_last=3), "TAS", path)
        wb = load_workbook(path, data_only=True)
        rows = {r[0]: r[1:] for r in wb["Summary"].iter_rows(min_row=5, values_only=True) if r[0]}
        for period in config.ROLLING_PERIODS:
            assert f"{period}-Year Average" in rows, f"period {period} row missing"
        assert rows["10-Year Average"][0] == 50.0
        assert rows["15-Year Average"][0] == "n/a" and "150 months" in rows["15-Year Average"][4]
        assert rows["20-Year Average"][0] == "n/a"
        text = " ".join(str(c) for r in wb["Summary"].iter_rows(values_only=True) for c in r if c)
        assert "Jun 2026 dollars" in text and "3 month" in text, text
        head = [c.value for c in wb["Monthly Data"][1]]
        assert head[-1] == "CPI Estimated"
        flags = [r[6] for r in wb["Monthly Data"].iter_rows(min_row=2, values_only=True)]
        assert flags.count("Yes") == 3 and flags[-1] == "Yes" and flags[0] in (None, "")


# ----------------------------------------------------------------------------- run(): full refresh + partial month
def _run(full_refresh, months, calls, partial=None, summary_rows=None, now=datetime(2026, 4, 20),
         fail_month=None, fail_mode="partial"):
    """Drive main.run() against fakes. Returns the summary written.

    fail_month makes that month fail in EVERY region: "partial" (incomplete file), "404" (empty frame,
    as download_month returns for a 404) or "error" (download_month raises).
    """
    with tempfile.TemporaryDirectory() as d:
        saved = {}
        patches = {
            (main, "PROJECT_ROOT"): Path(d),
            (main, "get_latest_available_month"): lambda: months[-1],
            (main, "generate_all_workbooks"): lambda *a, **k: None,
            (main, "get_cpi_lookup"): lambda cache: (
                pd.DataFrame({"year_month": [f"{y}-{m:02d}" for y, m in months], "cpi_index": 100.0}), 100.0),
            (config, "START_DATE"): datetime(*months[0], 1),
            (config, "REGION_START_DATES"): {r: datetime(*months[0], 1) for r in config.REGIONS},
            (config, "nem_now"): lambda: now,
        }

        def fake_download(year, month, region, cache_dir, force=False):
            calls.append((region, f"{year}-{month:02d}", force))
            if fail_month == f"{year}-{month:02d}":
                if fail_mode == "error":
                    raise RuntimeError("Failed to download after 3 attempts")
                if fail_mode == "404":
                    return pd.DataFrame(columns=download.EXPECTED_COLUMNS)
                return _month_frame(year, month, drop=500)
            if partial == (region, f"{year}-{month:02d}"):
                return _month_frame(year, month, drop=500)
            return _month_frame(year, month)

        patches[(main, "download_month")] = fake_download
        for (obj, name), val in patches.items():
            saved[(obj, name)] = getattr(obj, name)
            setattr(obj, name, val)
        try:
            if summary_rows is not None:
                (Path(d) / "outputs").mkdir()
                summary_rows.to_csv(Path(d) / "outputs" / "summary.csv", index=False)
            main.run(full_refresh=full_refresh, months_back=1)
            return pd.read_csv(Path(d) / "outputs" / "summary.csv")
        finally:
            for (obj, name), val in saved.items():
                setattr(obj, name, val)


MONTHS = [(2026, 1), (2026, 2), (2026, 3)]


def test_full_refresh_redownloads_everything():
    calls = []
    out = _run(True, MONTHS, calls)
    assert len(out) == 3 * len(config.REGIONS)
    assert all(force for _, _, force in calls), "--full-refresh must bypass the local cache"
    assert len(calls) == len(out)


def test_the_in_progress_month_is_excluded_using_nem_time():
    calls = []
    out = _run(True, MONTHS, calls, now=datetime(2026, 3, 20))   # March is still running
    assert sorted(out.year_month.unique()) == ["2026-01", "2026-02"]


def test_a_partial_month_is_not_published_and_the_rest_still_is():
    calls = []
    out = _run(True, MONTHS, calls, partial=("QLD1", "2026-03"))
    assert len(out) == 3 * len(config.REGIONS) - 1
    assert out[(out.region == "QLD1") & (out.year_month == "2026-03")].empty


def test_newest_month_failing_in_every_region_fails_the_run():
    # The old loop logged and skipped each failure, so a month no region produced was a green run.
    for mode in ("partial", "404", "error"):
        try:
            _run(True, MONTHS, [], fail_month="2026-03", fail_mode=mode)
        except RuntimeError as exc:
            assert "2026-03" in str(exc), exc
        else:
            raise AssertionError(f"newest month failing everywhere ({mode}) must fail the run")


def test_newest_month_already_published_survives_a_failed_redownload():
    # 2026-03 is already in summary.csv; its mutable-window re-download failing keeps the old row.
    calls = []
    first = _run(True, MONTHS, calls)
    out = _run(False, MONTHS, calls, summary_rows=first, fail_month="2026-03", fail_mode="error")
    assert len(out) == 3 * len(config.REGIONS)


def test_an_older_month_failing_everywhere_does_not_fail_the_run():
    # Only the newest month is checked here; a gap further back fails the validator's contiguity check.
    out = _run(True, MONTHS, [], fail_month="2026-02")
    assert "2026-02" not in set(out.year_month)


def test_latest_required_month_bound():
    # Aug ends 00:00 1 Sep; with a 35-day tolerance it is due from 00:00 6 Oct.
    assert config.AEMO_MAX_MONTH_LAG_DAYS == 35
    assert config.latest_required_month(datetime(2026, 10, 5, 23, 59)) == "2026-07"
    assert config.latest_required_month(datetime(2026, 10, 6, 0, 0)) == "2026-08"
    # A run on the 1st never requires the month that just ended, nor the one before it.
    assert config.latest_required_month(datetime(2026, 11, 1, 10, 38)) == "2026-08"
    assert config.latest_required_month(datetime(2027, 1, 7)) == "2026-11"    # year boundary
    # 28 days of AEMO lag (the most seen) stays inside the tolerance.
    assert config.latest_required_month(datetime(2026, 10, 29)) < "2026-09"


if __name__ == "__main__":
    tests = [v for k, v in sorted(globals().items()) if k.startswith("test_")]
    for t in tests:
        t()
        print(f"  PASS  {t.__name__}")
    print(f"\n{len(tests)} pipeline logic tests passed")
