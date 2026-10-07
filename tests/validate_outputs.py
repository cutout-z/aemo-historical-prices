"""Post-pipeline validation for AEMO Historical Prices.

Checks summary.csv and the Excel workbooks for integrity before anything is committed. Exits
non-zero on any failure. The checks are exact wherever the right answer is deterministic:

* every month holds EXACTLY days x 48 (30-min, before Oct 2021) or days x 288 (5-min) intervals, and
  exactly weekdays x 30 / x 180 peak intervals -- no exemptions, including the latest month;
* each region is contiguous from its configured start to a common end month, and that month is not
  older than config.latest_required_month() (a month that ended AEMO_MAX_MONTH_LAG_DAYS ago is due);
* the CPI columns obey their own invariants (flag is a suffix, real == nominal when flagged, one
  CPI ratio per month across regions and across the RRP / peak columns, cpi_base -- when present --
  is one month and the flag covers exactly the months after it);
* the carbon flag covers exactly Jul 2012 - Jun 2014;
* every workbook agrees with summary.csv.
"""

import calendar
import sys
from pathlib import Path

import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from src import config  # noqa: E402

OUTPUTS_DIR = Path(__file__).parent.parent / "outputs"
REGIONS = list(config.REGIONS)
REGION_NAMES = dict(config.REGION_NAMES)
FORMAT_CHANGE = config.FORMAT_CHANGE_DATE.strftime("%Y-%m")  # NEM moved from 30-min to 5-min
CARBON_FIRST, CARBON_LAST = "2012-07", "2014-06"
PRICE_COLS = ["rrp_nominal", "peak_rrp_nominal", "rrp_real", "peak_rrp_real"]
# The market price floor is -$1,000/MWh: no monthly mean can be below it.
MARKET_FLOOR = -1000.0

errors = []


def check(condition, msg):
    if not condition:
        errors.append(msg)
        print(f"  FAIL: {msg}")
    return condition


def _weekdays_in_month(year_month: str) -> int:
    y, m = int(year_month[:4]), int(year_month[5:])
    return sum(1 for d in range(1, calendar.monthrange(y, m)[1] + 1) if calendar.weekday(y, m, d) < 5)


def _expected_intervals(year_month: str) -> int:
    return config.expected_interval_count(int(year_month[:4]), int(year_month[5:]))


def _expected_peak(year_month: str) -> int:
    per_day = 180 if year_month >= FORMAT_CHANGE else 30   # 15 peak hours x 12 (5-min) or x 2 (30-min)
    return _weekdays_in_month(year_month) * per_day


def _month_range(first: str, last: str) -> list[str]:
    return [p.strftime("%Y-%m") for p in pd.period_range(first, last, freq="M")]


def validate():
    summary_path = OUTPUTS_DIR / "summary.csv"
    check(summary_path.exists(), "summary.csv does not exist")
    if not summary_path.exists():
        return

    df = pd.read_csv(summary_path)
    print(f"summary.csv: {len(df)} rows")

    # --- Structure ---
    check(len(df) > 0, "summary.csv is empty")
    required_cols = ["region", "year_month", *PRICE_COLS, "total_intervals", "peak_intervals",
                     "carbon_flag", "cpi_estimated"]
    missing = [c for c in required_cols if c not in df.columns]
    for col in missing:
        check(False, f"Missing column: {col}")
    if missing or df.empty:
        return

    # --- All 5 regions present ---
    regions_present = set(df["region"].unique())
    for r in REGIONS:
        check(r in regions_present, f"Region {r} missing from summary.csv")

    # --- No duplicate region/month ---
    dupes = df.duplicated(subset=["region", "year_month"], keep=False)
    check(dupes.sum() == 0, f"{dupes.sum()} duplicate region/month rows")

    # --- Values present and finite ---
    for col in PRICE_COLS:
        vals = pd.to_numeric(df[col], errors="coerce")
        check(vals.notna().all() and vals.abs().lt(float("inf")).all(),
              f"{col} has {int(vals.isna().sum())} missing / non-numeric value(s)")
        check(vals.min() >= MARKET_FLOOR, f"{col} is below the market floor (min={vals.min():.2f})")
    check((df["rrp_real"] > 0).all(), "rrp_real has non-positive values")

    # --- Contiguity: each region runs unbroken from its start to ONE common end month ---
    ends = df.groupby("region")["year_month"].max()
    check(ends.nunique() == 1, f"regions end on different months: {ends.to_dict()}")
    common_end = ends.max()
    required = config.latest_required_month(config.nem_now())
    check(common_end >= required,
          f"newest month is {common_end}, but {required} ended at least {config.AEMO_MAX_MONTH_LAG_DAYS} "
          f"days ago and is missing (AEMO not publishing, or every region's download failing?)")
    for region in REGIONS:
        have = sorted(df.loc[df["region"] == region, "year_month"])
        if not have:
            continue
        start = config.REGION_START_DATES[region].strftime("%Y-%m")
        want = _month_range(start, common_end)
        check(have == want,
              f"{region} is not contiguous {start}..{common_end}: "
              f"{len(set(want) - set(have))} missing, {len(set(have) - set(want))} unexpected")

    # --- Interval counts: exact for every month, no exemptions ---
    exp_total = df["year_month"].map(_expected_intervals)
    bad_total = df[df["total_intervals"] != exp_total]
    check(bad_total.empty,
          f"{len(bad_total)} rows have an incomplete / wrong interval count, e.g. "
          f"{bad_total[['region', 'year_month', 'total_intervals']].head(3).values.tolist()}")

    exp_peak = df["year_month"].map(_expected_peak)
    bad_peak = df[df["peak_intervals"] != exp_peak]
    check(bad_peak.empty,
          f"{len(bad_peak)} rows have the wrong number of peak intervals, e.g. "
          f"{bad_peak[['region', 'year_month', 'peak_intervals']].head(3).values.tolist()}")

    # --- Carbon flag: exactly Jul 2012 - Jun 2014 ---
    flag = df["carbon_flag"].astype(str).str.lower().eq("true")
    want_flag = (df["year_month"] >= CARBON_FIRST) & (df["year_month"] <= CARBON_LAST)
    check((flag == want_flag).all(), f"carbon_flag differs from {CARBON_FIRST}..{CARBON_LAST} on "
                                    f"{int((flag != want_flag).sum())} rows")

    # --- CPI invariants ---
    est = df["cpi_estimated"].astype(str).str.lower().eq("true")
    for region in REGIONS:
        r = est[df["region"] == region].values[
            df.loc[df["region"] == region, "year_month"].argsort().values]
        # once True, always True: months without CPI are only ever the most recent ones
        check(not any(a and not b for a, b in zip(r, r[1:])),
              f"{region}: cpi_estimated is not a suffix of the series")
    est_sets = df[est].groupby("region")["year_month"].apply(frozenset)
    check(len(set(est_sets)) <= 1, "regions disagree on which months are cpi_estimated")
    flagged = df[est]
    check((flagged["rrp_real"] == flagged["rrp_nominal"]).all()
          and (flagged["peak_rrp_real"] == flagged["peak_rrp_nominal"]).all(),
          "a cpi_estimated month has real != nominal (it must carry ratio 1)")
    check(est.mean() < 0.05, f"{est.mean():.1%} of rows are cpi_estimated (CPI fetch may have failed)")
    # cpi_base (newest CPI quarter, the real-dollar base) is optional so older files still validate.
    if "cpi_base" in df.columns:
        bases = df["cpi_base"].dropna().astype(str).unique()
        if check(df["cpi_base"].notna().all() and len(bases) == 1,
                 f"cpi_base must be one value on every row, got {list(bases)[:3]}"):
            base = bases[0]
            check(not df.loc[est, "year_month"].le(base).any()
                  and not df.loc[~est, "year_month"].gt(base).any(),
                  f"cpi_estimated must flag exactly the months after cpi_base {base}")

    ok = df[~est & (df["rrp_nominal"] >= 5)].copy()
    ok["ratio"] = ok["rrp_real"] / ok["rrp_nominal"]
    spread = ok.groupby("year_month")["ratio"].agg(lambda x: x.max() - x.min())
    check((spread <= 0.003).all(),
          f"the CPI ratio differs between regions in {int((spread > 0.003).sum())} month(s), "
          f"e.g. {spread[spread > 0.003].head(3).round(4).to_dict()}")
    pk = df[~est & (df["peak_rrp_nominal"] >= 20) & (df["rrp_nominal"] >= 5)].copy()
    drift = (pk["peak_rrp_real"] / pk["peak_rrp_nominal"] - pk["rrp_real"] / pk["rrp_nominal"]).abs()
    check((drift <= 0.003).all(),
          f"peak and all-hours prices use different CPI ratios on {int((drift > 0.003).sum())} rows")

    # --- Workbooks exist and agree with summary.csv ---
    from openpyxl import load_workbook

    for region_id, name in REGION_NAMES.items():
        xlsx_path = OUTPUTS_DIR / f"{name}_historical_prices.xlsx"
        if not check(xlsx_path.exists(), f"{xlsx_path.name} does not exist"):
            continue
        want = df[df["region"] == region_id].sort_values("year_month")
        wb = load_workbook(xlsx_path, read_only=True, data_only=True)
        check(set(wb.sheetnames) == {"Summary", "Monthly Data", "Heatmap"},
              f"{xlsx_path.name}: unexpected sheets {wb.sheetnames}")
        if "Monthly Data" in wb.sheetnames:
            rows = [r for r in wb["Monthly Data"].iter_rows(min_row=2, values_only=True) if r[0]]
            check(len(rows) == len(want), f"{xlsx_path.name}: {len(rows)} data rows vs {len(want)} in summary.csv")
            if len(rows) == len(want):
                got = pd.DataFrame(rows).iloc[:, 1:5].astype(float).round(2).values
                check((got == want[PRICE_COLS].round(2).values).all(),
                      f"{xlsx_path.name}: prices differ from summary.csv")
        wb.close()

    all_states_path = OUTPUTS_DIR / "All_States_historical_prices.xlsx"
    if check(all_states_path.exists(), "All_States_historical_prices.xlsx does not exist"):
        wb = load_workbook(all_states_path, read_only=True, data_only=True)
        check(wb.sheetnames == [REGION_NAMES[r] for r in REGIONS],
              f"All_States workbook sheets are {wb.sheetnames}")
        for region_id, name in REGION_NAMES.items():
            if name in wb.sheetnames:
                n = sum(1 for r in wb[name].iter_rows(min_row=2, values_only=True) if r[0])
                want_n = int((df["region"] == region_id).sum())
                check(n == want_n, f"All_States {name}: {n} rows vs {want_n} in summary.csv")
        wb.close()


if __name__ == "__main__":
    print("Validating AEMO Historical Prices outputs...")
    validate()
    if errors:
        print(f"\n{len(errors)} validation error(s) found — aborting.")
        sys.exit(1)
    else:
        print("\nAll validations passed.")
